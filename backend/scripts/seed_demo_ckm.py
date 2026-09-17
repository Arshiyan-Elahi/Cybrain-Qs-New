"""Idempotently seed a self-contained CKM demo tenant for local manual testing."""
import io
from datetime import datetime, timezone

from docx import Document as DocxDocument
from sqlalchemy import select

from app.core.database import get_session_factory
from app.core.security import hash_password, verify_password
from app.modules.auth.models import User, UserCompanyAccess
from app.modules.companies.models import Company, CompanyRegulation
from app.modules.documents.ingestion import IngestionService
from app.modules.documents.models import Document, DocumentChunk
from app.modules.knowledge.models import KnowledgeObject
from app.shared.enums import KnowledgeStatus

DEMO_EMAIL = "demo@cybrain.local"
DEMO_PASSWORD = "CybrainDemo123!"
COMPANY_NAME = "Cybrain Demo Pharma"

SOPS = {
    "SOP-001-Document-Control.docx": [
        ("1. Purpose", "This SOP defines GMP document control for Cybrain Demo Pharma."),
        ("2. Responsibilities", "The Quality Assurance Manager approves controlled documents. The Document Owner drafts and reviews revisions."),
        ("3. Procedure", "The Document Owner assigns an SOP number. Quality Assurance reviews the draft against EU GMP and Annex 11. Approved versions are released; obsolete versions are archived."),
        ("4. Business Rules", "Only the current approved version may be used. Training must be completed before the effective date."),
    ],
    "SOP-002-Deviation-Management.docx": [
        ("1. Scope", "This procedure applies to GMP deviations in manufacturing and quality operations."),
        ("2. Terminology", "A deviation is an unplanned departure from an approved instruction or expected result. CAPA means corrective and preventive action."),
        ("3. Workflow", "The Reporter records the event within one business day. The Investigator determines root cause. Quality Assurance assesses impact and approves closure."),
        ("4. Regulations", "Records must be attributable, legible, contemporaneous, original and accurate in accordance with EU GMP and ICH Q10."),
    ],
    "SOP-003-Training.docx": [
        ("1. Purpose", "This SOP establishes a role-based training process."),
        ("2. Writing Standard", "Instructions use active voice, mandatory wording with shall, and numbered steps. Abbreviations are defined on first use."),
        ("3. Responsibilities", "Department Managers assign curricula. Employees complete training. Quality Assurance monitors overdue training."),
        ("4. Procedure", "Assign training after approval, notify the employee, record completion, and retain the training record for inspection."),
    ],
}

FIXTURES = [
    ("terminology", "CAPA — corrective and preventive action"),
    ("writing_style", "Use active voice, mandatory wording, and numbered steps"),
    ("document_structure", "Purpose, responsibilities, procedure, and records sections"),
    ("regulation", "EU GMP and ICH Q10"),
    ("role", "Quality Assurance Manager approves controlled documents"),
    ("workflow", "Report, investigate, assess impact, and approve closure"),
    ("business_rule", "Only the current approved SOP version may be used"),
]


def _docx_bytes(sections: list[tuple[str, str]]) -> bytes:
    document = DocxDocument()
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def main() -> None:
    session = get_session_factory()()
    try:
        user = session.scalar(select(User).where(User.email == DEMO_EMAIL))
        if user is None:
            user = User(email=DEMO_EMAIL, hashed_password=hash_password(DEMO_PASSWORD), full_name="Demo QA Reviewer")
            session.add(user)
            session.flush()
        else:
            # A repeatable dev seed must repair stale fixture credentials and a
            # disabled fixture account without changing any non-demo user.
            if not verify_password(DEMO_PASSWORD, user.hashed_password):
                user.hashed_password = hash_password(DEMO_PASSWORD)
            user.is_active = True
            user.full_name = "Demo QA Reviewer"
            session.flush()

        company = session.scalar(select(Company).where(Company.name == COMPANY_NAME))
        if company is None:
            company = Company(name=COMPANY_NAME, industry_key="pharma", location_key="berlin-de", primary_language_key="en", sop_count=3)
            company.regulations = [
                CompanyRegulation(regulation_id=value, position=index)
                for index, value in enumerate(["eu-gmp", "annex-11", "ich-q10"])
            ]
            session.add(company)
            session.flush()
        if session.scalar(select(UserCompanyAccess).where(
            UserCompanyAccess.user_id == user.id, UserCompanyAccess.company_id == company.id
        )) is None:
            session.add(UserCompanyAccess(user_id=user.id, company_id=company.id, role="owner"))
            session.flush()

        ingestion = IngestionService(session)
        for filename, sections in SOPS.items():
            existing = session.scalar(select(Document).where(
                Document.company_id == company.id, Document.filename == filename
            ))
            if existing is None:
                ingestion.ingest(company_id=company.id, filename=filename,
                                 content=_docx_bytes(sections), uploaded_by=user.id, embed=False)

        chunks = list(session.scalars(select(DocumentChunk).where(
            DocumentChunk.company_id == company.id
        ).order_by(DocumentChunk.document_id, DocumentChunk.chunk_order)).all())
        existing_types = set(session.scalars(select(KnowledgeObject.type).where(
            KnowledgeObject.company_id == company.id,
            KnowledgeObject.extraction_method == "dev-seed",
        )).all())
        for index, (kind, label) in enumerate(FIXTURES):
            if kind in existing_types or not chunks:
                continue
            chunk = chunks[index % len(chunks)]
            session.add(KnowledgeObject(
                company_id=company.id, type=kind, tier=chunk.tier,
                status=KnowledgeStatus.PROPOSED, label=label,
                payload={"fixture": True}, source_document_id=chunk.document_id,
                source_chunk_id=chunk.id, source_location=chunk.location,
                extraction_method="dev-seed", extracted_at=datetime.now(timezone.utc),
            ))
        session.commit()
        print(f"Seeded {COMPANY_NAME}: {len(SOPS)} documents, {len(FIXTURES)} proposed object types.")
        print(f"Login: {DEMO_EMAIL} / {DEMO_PASSWORD}")
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


if __name__ == "__main__":
    main()
