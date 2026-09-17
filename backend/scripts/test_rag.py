"""
End-to-end RAG test against the configured local endpoint and a live database.

Ingests real DOCX documents for two different companies, embeds them with the
configured model, and checks that retrieval is relevant, tier-aware and — most
importantly — cannot cross the company boundary.

    python scripts/test_rag.py
"""

import io
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import docx  # noqa: E402

from app.core.database import get_session_factory  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.integrations.llm.factory import get_llm_provider  # noqa: E402
from app.modules.auth.models import User, UserCompanyAccess  # noqa: E402
from app.modules.companies.models import Company  # noqa: E402
from app.modules.documents.ingestion import IngestionService  # noqa: E402
from app.modules.knowledge.retrieval import RetrievalService  # noqa: E402
from app.shared.enums import DocumentStatus, KnowledgeTier  # noqa: E402

failures: list[str] = []


def check(label: str, actual, expected) -> None:
    ok = actual == expected
    print(f"{'PASS' if ok else 'FAIL'}  {label}  (got {actual!r}, want {expected!r})")
    if not ok:
        failures.append(label)


def check_true(label: str, value) -> None:
    check(label, bool(value), True)


def make_docx(title: str, sections: list[tuple[str, str]]) -> bytes:
    document = docx.Document()
    document.add_heading(title, level=1)
    for heading, body in sections:
        document.add_heading(heading, level=2)
        document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


db = get_session_factory()()
llm = get_llm_provider()
tag = uuid.uuid4().hex[:8]

print(f"provider: {llm.name} | chat: {llm.chat_model} | embed: {llm.embedding_model}\n")

# --- fixtures -------------------------------------------------------------
user = User(
    email=f"rag-{tag}@example.com", hashed_password=hash_password("x" * 12), full_name="RAG Tester"
)
pharma = Company(name=f"Pharma {tag}", industry_key="pharma", location_key="vienna-at")
rival = Company(name=f"Rival {tag}", industry_key="pharma", location_key="berlin-de")
db.add_all([user, pharma, rival])
db.flush()
db.add_all(
    [
        UserCompanyAccess(user_id=user.id, company_id=pharma.id, role="owner"),
        UserCompanyAccess(user_id=user.id, company_id=rival.id, role="owner"),
    ]
)
db.commit()

ingestion = IngestionService(db, llm)
retrieval = RetrievalService(db, llm)

print("--- ingestion + embedding (real model) ---")
pharma_doc = ingestion.ingest(
    company_id=pharma.id,
    filename="SOP-014-Lieferantenqualifizierung.docx",
    content=make_docx(
        "SOP-014 Lieferantenqualifizierung",
        [
            ("1. Zweck", "Diese SOP regelt die Qualifizierung und Bewertung externer Lieferanten."),
            ("2. Verantwortlichkeiten", "Der QA Manager gibt die Lieferantenqualifizierung frei. Der Einkauf waehlt Lieferanten aus."),
            ("3. Auditplanung", "Lieferantenaudits werden jaehrlich geplant und dokumentiert."),
        ],
    ),
    uploaded_by=user.id,
)
check("document processed", pharma_doc.status, DocumentStatus.PROCESSED)
check_true("chunks created", len(pharma_doc.chunks) >= 3)
check_true("all chunks embedded", all(c.embedding is not None for c in pharma_doc.chunks))
check("embedding dimension stored", len(pharma_doc.chunks[0].embedding), llm.embedding_dimensions)
check("embedding model recorded", pharma_doc.chunks[0].embedding_model, llm.embedding_model)

# A second company with deliberately distinctive content.
rival_doc = ingestion.ingest(
    company_id=rival.id,
    filename="SOP-900-Kalibrierung.docx",
    content=make_docx(
        "SOP-900 Kalibrierung von Messmitteln",
        [
            ("1. Zweck", "Diese SOP beschreibt die Kalibrierung von Waagen und Thermometern."),
            ("2. Intervalle", "Messmittel werden halbjaehrlich kalibriert und protokolliert."),
        ],
    ),
    uploaded_by=user.id,
)
check("second document processed", rival_doc.status, DocumentStatus.PROCESSED)

print("\n--- retrieval relevance ---")
hits = retrieval.search(company_id=pharma.id, query="Wer gibt die Lieferantenqualifizierung frei?", limit=3)
check_true("hits returned", len(hits) > 0)
check_true("top hit mentions QA Manager", "QA Manager" in hits[0].text)
check_true("similarity is sane", 0.0 < hits[0].similarity <= 1.0)
check_true("provenance: filename", hits[0].document_filename.startswith("SOP-014"))
check_true("provenance: location has heading path", ">" in hits[0].location or hits[0].heading_path)
check("tier attached", hits[0].tier, KnowledgeTier.COMPANY)

hits_audit = retrieval.search(company_id=pharma.id, query="Wie oft finden Audits statt?", limit=3)
check_true(
    "audit query retrieves the audit section",
    any("audit" in h.text.lower() or "audit" in " ".join(h.heading_path).lower() for h in hits_audit),
)
# The heading path is embedded with the body, so the audit section should now
# outrank the others for an audit question rather than tie with them.
check_true(
    "audit section ranks first",
    "audit" in (hits_audit[0].text + " ".join(hits_audit[0].heading_path)).lower(),
)

print("\n--- company isolation (the one that matters) ---")
cross = retrieval.search(company_id=rival.id, query="Wer gibt die Lieferantenqualifizierung frei?", limit=5)
check_true("rival gets its own hits only", all(h.document_id == rival_doc.id for h in cross))
check_true("rival never sees pharma text", all("Lieferantenqualifizierung" not in h.text for h in cross))

pharma_ids = {h.document_id for h in retrieval.search(company_id=pharma.id, query="Kalibrierung von Waagen", limit=5)}
check("pharma never sees rival document", rival_doc.id in pharma_ids, False)

print("\n--- tier filtering ---")
industry_doc = ingestion.ingest(
    company_id=pharma.id,
    filename="ICH-Q10-guidance.docx",
    content=make_docx("ICH Q10 Guidance", [("Supplier control", "Sector guidance on supplier qualification expectations.")]),
    tier=KnowledgeTier.INDUSTRY,
)
check("industry doc ingested", industry_doc.status, DocumentStatus.PROCESSED)

company_only = retrieval.search(company_id=pharma.id, query="supplier qualification", limit=10, tiers=[KnowledgeTier.COMPANY])
check_true("company-only excludes industry", all(h.tier == KnowledgeTier.COMPANY for h in company_only))

both = retrieval.search(company_id=pharma.id, query="supplier qualification", limit=10)
check_true("unfiltered includes industry tier", any(h.tier == KnowledgeTier.INDUSTRY for h in both))

context = retrieval.build_context(both[:3])
check_true("context labels every block with its tier", context.count("tier=") == len(both[:3]))
check_true("industry blocks marked as not the client's", "not this client's" in context or all(h.tier == KnowledgeTier.COMPANY for h in both[:3]))

print("\n--- cleanup ---")
for company in (pharma, rival):
    db.delete(company)
db.delete(user)
db.commit()
db.close()

print(f"\n{'ALL PASSED' if not failures else str(len(failures)) + ' FAILED: ' + ', '.join(failures)}")
sys.exit(1 if failures else 0)
