"""
DOCX extraction with multi-signal heading and structure detection.

Does not rely on Heading styles alone. Signals include style name, numbering,
font size, bold, indentation, spacing, document position, surrounding blocks
and table context.
"""

from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass, field
from typing import Any

from app.processing.errors import UnsupportedDocumentError
from app.processing.schema import ExtractionQA, NormalizedBlock, NormalizedDocument

_NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+(\S.*)$")
_LIST_ITEM = re.compile(r"^\s*(?:[-•*•]|\(?[a-z0-9]{1,3}[.)])\s+\S", re.I)
_APPENDIX = re.compile(r"^\s*appendix\s+([ivxlcdm\d]+)\b", re.I)
_REVISION_HINT = re.compile(
    r"\b(revision|change\s*history|document\s*history|amendment\s*history|"
    r"versionshistorie|änderungsHistorie|änderungshistorie)\b",
    re.I,
)
_APPROVAL_HINT = re.compile(
    r"\b(signature|signed|approval|approved\s+by|freigabe|unterschrift|"
    r"reviewed\s+by|authori[sz]ed)\b",
    re.I,
)


@dataclass
class _RawPara:
    text: str
    style_name: str
    is_heading_style: bool
    heading_style_level: int | None
    has_numbering: bool
    font_size_pt: float | None
    bold: bool | None
    indent_pt: float
    space_before_pt: float
    space_after_pt: float
    in_table: bool
    order: int


@dataclass
class _RawTable:
    rows: list[list[str]]
    order: int
    texts: str = ""


def extract_docx(content: bytes) -> NormalizedDocument:
    import docx
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        raise UnsupportedDocumentError(f"Could not read DOCX: {exc}") from exc

    raw_items: list[tuple[str, Any]] = []
    order = 0
    font_sizes: list[float] = []

    def _para_signals(paragraph: Paragraph) -> _RawPara:
        nonlocal order
        style_name = (paragraph.style.name if paragraph.style else "") or ""
        style_l = style_name.lower()
        heading_match = re.match(r"heading\s*(\d+)", style_l)
        custom_heading = any(
            token in style_l for token in ("überschrift", "heading", "title", "titel")
        )
        level = int(heading_match.group(1)) if heading_match else (1 if "title" in style_l else None)

        has_numbering = False
        pPr = paragraph._element.pPr
        if pPr is not None and pPr.numPr is not None:
            has_numbering = True

        sizes: list[float] = []
        bold_flags: list[bool] = []
        for run in paragraph.runs:
            if run.bold is not None:
                bold_flags.append(bool(run.bold))
            if run.font.size is not None:
                sizes.append(run.font.size.pt)
            elif run.style and run.style.font and run.style.font.size:
                sizes.append(run.style.font.size.pt)
        font_size = max(sizes) if sizes else None
        if font_size is not None:
            font_sizes.append(font_size)
        bold = None
        if bold_flags:
            bold = all(bold_flags) or (len(bold_flags) == 1 and bold_flags[0])

        indent = 0.0
        if paragraph.paragraph_format.left_indent is not None:
            indent = float(paragraph.paragraph_format.left_indent.pt)
        space_before = (
            float(paragraph.paragraph_format.space_before.pt)
            if paragraph.paragraph_format.space_before
            else 0.0
        )
        space_after = (
            float(paragraph.paragraph_format.space_after.pt)
            if paragraph.paragraph_format.space_after
            else 0.0
        )
        _ = (WD_PARAGRAPH_ALIGNMENT, qn)  # imported for OOXML introspection stability

        item = _RawPara(
            text=paragraph.text.strip(),
            style_name=style_name,
            is_heading_style=bool(heading_match or custom_heading),
            heading_style_level=level,
            has_numbering=has_numbering,
            font_size_pt=font_size,
            bold=bold,
            indent_pt=indent,
            space_before_pt=space_before,
            space_after_pt=space_after,
            in_table=False,
            order=order,
        )
        order += 1
        return item

    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            raw_items.append(("p", _para_signals(paragraph)))
        elif tag == "tbl":
            table = Table(child, document)
            cells = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if any(any(c for c in row) for row in cells):
                joined = "\n".join(" | ".join(row) for row in cells if any(row))
                raw_items.append(
                    ("t", _RawTable(rows=cells, order=order, texts=joined))
                )
                order += 1

    # Headers / footers — retained for audit, marked non-semantic later.
    header_footer_texts: list[str] = []
    for section in document.sections:
        for hf in (section.header, section.footer):
            for paragraph in hf.paragraphs:
                text = paragraph.text.strip()
                if text:
                    header_footer_texts.append(text)

    median_size = _median(font_sizes) or 11.0
    blocks: list[NormalizedBlock] = []
    block_order = 0
    seen_heading = False
    appendix_active = False

    for kind, item in raw_items:
        if kind == "t":
            table: _RawTable = item
            block_type, semantic = _classify_table(table.texts, table.rows, seen_heading)
            blocks.append(
                NormalizedBlock(
                    id=_bid(),
                    type=block_type,
                    text=table.texts,
                    raw_text=table.texts,
                    order=block_order,
                    semantic=semantic,
                    table={"rows": table.rows},
                    style={"source": "docx_table"},
                    provenance={"parser": "docx", "source_order": table.order},
                )
            )
            block_order += 1
            appendix_active = False
            continue

        para: _RawPara = item
        classified = _classify_paragraph(
            para,
            median_size=median_size,
            seen_heading=seen_heading,
            appendix_active=appendix_active,
        )
        if classified["type"] in ("heading", "appendix"):
            seen_heading = True
            appendix_active = bool(
                classified["type"] == "appendix" or _APPENDIX.match(para.text)
            )
        elif appendix_active:
            appendix_active = False
        blocks.append(
            NormalizedBlock(
                id=_bid(),
                type=classified["type"],
                text=para.text,
                raw_text=para.text,
                order=block_order,
                heading_level=classified.get("level"),
                semantic=classified.get("semantic", True),
                style={
                    "style_name": para.style_name,
                    "font_size_pt": para.font_size_pt,
                    "bold": para.bold,
                    "indent_pt": para.indent_pt,
                    "space_before_pt": para.space_before_pt,
                    "space_after_pt": para.space_after_pt,
                    "has_numbering": para.has_numbering,
                    "signals": classified.get("signals", []),
                },
                provenance={"parser": "docx", "source_order": para.order},
            )
        )
        block_order += 1

    for text in _unique_preserve(header_footer_texts):
        blocks.append(
            NormalizedBlock(
                id=_bid(),
                type="header_footer",
                text=text,
                raw_text=text,
                order=block_order,
                semantic=False,
                style={"source": "docx_header_footer"},
                provenance={"parser": "docx", "furniture": True},
            )
        )
        block_order += 1

    core = document.core_properties
    metadata = {
        "title": core.title or "",
        "author": core.author or "",
        "page_count": 0,
        "source_format": "docx",
    }

    warnings = [] if blocks else ["The document contains no readable text."]
    return NormalizedDocument(
        metadata=metadata,
        pages=[],
        blocks=blocks,
        warnings=warnings,
        extraction_quality=ExtractionQA(status="PASS", diagnostics=[]),
        source_format="docx",
        parser="python-docx",
    )


def _classify_paragraph(
    para: _RawPara,
    *,
    median_size: float,
    seen_heading: bool,
    appendix_active: bool = False,
) -> dict[str, Any]:
    text = para.text
    signals: list[str] = []
    score = 0.0

    if para.is_heading_style:
        score += 4
        signals.append("style")
    if para.heading_style_level:
        signals.append(f"style_level:{para.heading_style_level}")

    numbered = _NUMBERED_HEADING.match(text)
    if numbered and len(text) <= 160:
        score += 3
        signals.append("numbering_text")

    if para.has_numbering and len(text) <= 160 and not _LIST_ITEM.match(text):
        score += 2
        signals.append("numPr")

    if para.font_size_pt and para.font_size_pt >= median_size + 1.5:
        score += 2
        signals.append("font_size")
    if para.bold and len(text) <= 120:
        score += 1
        signals.append("bold")
    if para.space_before_pt >= 8 or para.space_after_pt >= 6:
        score += 1
        signals.append("spacing")
    if para.indent_pt <= 0 and len(text) <= 100:
        score += 0.5
        signals.append("indent")

    if _APPENDIX.match(text):
        score += 3
        signals.append("appendix_label")

    # Short title-ish lines without terminal punctuation.
    if 0 < len(text) <= 90 and not text.endswith((".", ":", ";", ",")):
        letters = [c for c in text if c.isalpha()]
        if letters and sum(c.isupper() for c in letters) / len(letters) > 0.65:
            score += 2
            signals.append("caps")

    style_l = para.style_name.lower()

    # Heading detection must win over list-item regex: "1. Zweck" looks like a
    # list marker but is a numbered section title in nearly every SOP.
    if score >= 3.5 or para.is_heading_style or (numbered and len(text) <= 120):
        level = para.heading_style_level
        if level is None and numbered:
            level = numbered.group(1).count(".") + 1
        if level is None and _APPENDIX.match(text):
            level = 1
        if level is None:
            level = 1 if score >= 5 else 2
        block_type = "appendix" if _APPENDIX.match(text) else "heading"
        return {"type": block_type, "level": level, "semantic": True, "signals": signals}

    # First short line after an Appendix label is treated as a sub-heading.
    if (
        appendix_active
        and len(text) <= 120
        and not text.endswith((".", ";", ","))
        and not _APPENDIX.match(text)
    ):
        return {
            "type": "heading",
            "level": 2,
            "semantic": True,
            "signals": ["appendix_followon"],
        }

    if style_l.startswith("list") or (
        _LIST_ITEM.match(text) and not _NUMBERED_HEADING.match(text)
    ):
        return {"type": "list", "semantic": True, "signals": ["list"]}

    if not seen_heading:
        return {"type": "document_metadata", "semantic": False, "signals": ["pre_heading"]}

    if (
        _REVISION_HINT.search(text)
        and len(text) <= 80
        and not text.endswith(".")
        and len(text.split()) <= 6
    ):
        return {"type": "heading", "level": 1, "semantic": True, "signals": ["revision_label"]}

    return {"type": "paragraph", "semantic": True, "signals": signals}


def _classify_table(
    texts: str, rows: list[list[str]], seen_heading: bool
) -> tuple[str, bool]:
    flat = texts.lower()
    if (not seen_heading or _looks_like_approval(rows, flat)) and _APPROVAL_HINT.search(flat):
        return "approval_signature", False
    if _REVISION_HINT.search(flat) or _looks_like_revision_table(rows):
        return "revision_history", True
    return "table", True


def _looks_like_approval(rows: list[list[str]], flat: str) -> bool:
    header = " ".join(rows[0]).lower() if rows else ""
    tokens = ("signature", "name", "date", "role", "approval", "freigabe", "unterschrift")
    return sum(1 for t in tokens if t in header or t in flat) >= 2


def _looks_like_revision_table(rows: list[list[str]]) -> bool:
    if not rows:
        return False
    header = " ".join(rows[0]).lower()
    tokens = ("rev", "revision", "version", "date", "author", "description", "change", "änderung")
    return sum(1 for t in tokens if t in header) >= 2


def _median(values: list[float]) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    mid = len(ordered) // 2
    if len(ordered) % 2:
        return ordered[mid]
    return (ordered[mid - 1] + ordered[mid]) / 2


def _unique_preserve(items: list[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for item in items:
        key = item.casefold()
        if key not in seen:
            seen.add(key)
            out.append(item)
    return out


def _bid() -> str:
    return uuid.uuid4().hex[:12]
