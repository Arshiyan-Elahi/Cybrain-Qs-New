"""Document parsing and chunking. Pure functions: no database, no server, no LLM."""

import io

import docx
import pytest
from pypdf import PdfWriter

from app.processing.chunker import chunk_document
from app.processing.extractor import UnsupportedDocumentError, extract


@pytest.fixture
def sop_docx() -> bytes:
    document = docx.Document()
    document.add_heading("SOP-014 Lieferantenqualifizierung", level=1)
    document.add_heading("1. Zweck", level=2)
    document.add_paragraph("Diese SOP beschreibt die Qualifizierung externer Lieferanten.")
    document.add_heading("2. Verantwortlichkeiten", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Rolle"
    table.cell(0, 1).text = "Verantwortung"
    table.cell(1, 0).text = "QA Manager"
    table.cell(1, 1).text = "Freigabe der Qualifizierung"
    document.add_heading("3. Ablauf", level=2)
    for step in ["Lieferantenauswahl", "Auditplanung", "Bewertung"]:
        document.add_paragraph(step, style="List Number")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestDocxExtraction:
    def test_structure_is_classified(self, sop_docx):
        kinds = [b.kind for b in extract(sop_docx, "sop.docx").blocks]
        assert kinds.count("heading") >= 4
        assert kinds.count("table") == 1
        assert kinds.count("list") >= 3

    def test_table_stays_whole(self, sop_docx):
        table = next(b for b in extract(sop_docx, "sop.docx").blocks if b.kind == "table")
        assert "QA Manager | Freigabe der Qualifizierung" in table.text


class TestChunking:
    def test_every_chunk_has_a_heading_path(self, sop_docx):
        chunks = chunk_document(extract(sop_docx, "sop.docx"))
        assert chunks and all(c.heading_path for c in chunks)

    def test_table_is_its_own_chunk(self, sop_docx):
        chunks = chunk_document(extract(sop_docx, "sop.docx"))
        table_chunk = next(c for c in chunks if "QA Manager" in c.text)
        assert table_chunk.heading_path[-1].startswith("2. Verantwortlichkeiten")

    def test_oversized_section_splits_but_keeps_path(self):
        document = docx.Document()
        document.add_heading("Big section", level=1)
        for i in range(60):
            document.add_paragraph(f"Absatz {i}: " + ("Lorem ipsum dolor sit amet. " * 6))
        buffer = io.BytesIO()
        document.save(buffer)

        chunks = chunk_document(extract(buffer.getvalue(), "big.docx"), max_chars=1200)
        assert len(chunks) > 1
        assert max(c.char_count for c in chunks) < 1500
        assert all(c.heading_path == ["Big section"] for c in chunks)


class TestFailureHandling:
    """Parsing problems must be surfaced, never silently swallowed."""

    def test_unsupported_type(self):
        with pytest.raises(UnsupportedDocumentError):
            extract(b"plain text", "notes.txt")

    def test_corrupt_pdf(self):
        with pytest.raises(UnsupportedDocumentError):
            extract(b"%PDF-1.4 truncated", "broken.pdf")

    def test_image_only_pdf_is_flagged_for_ocr(self):
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.add_blank_page(width=595, height=842)
        buffer = io.BytesIO()
        writer.write(buffer)

        result = extract(buffer.getvalue(), "scanned.pdf")
        assert result.is_probably_scanned
        assert any("OCR" in w for w in result.warnings)
        # No phantom blocks invented from blank pages.
        assert result.blocks == []


def test_sop_qa_001_normalizes_approval_and_appendix_hierarchy():
    """Regression shape for SOP-QA-001 v11 SOP Writing Numberering Archiving.docx."""
    document = docx.Document()
    approval = document.add_table(rows=2, cols=3)
    for index, value in enumerate(("Name", "Signature", "Date")):
        approval.cell(0, index).text = value
        approval.cell(1, index).text = "QA Approver"
    document.add_heading("1. References", level=1)
    document.add_paragraph("EU GMP Chapter 4")
    document.add_heading("5. Procedure", level=1)
    document.add_heading("5.1 General", level=2)
    document.add_paragraph("SOPs shall follow the approved template.")
    document.add_heading("5.2 Numbering of SOPs and WIs", level=2)
    document.add_paragraph("Assign one controlled identifier.")
    content_table = document.add_table(rows=2, cols=2)
    content_table.cell(0, 0).text = "Type"
    content_table.cell(0, 1).text = "Pattern"
    content_table.cell(1, 0).text = "SOP"
    content_table.cell(1, 1).text = "SOP-XX-000"
    document.add_heading("6. SOP Changes", level=1)
    document.add_paragraph("Changes require revision history.")
    document.add_paragraph("Appendix I")
    document.add_paragraph("SOP and WI Content")
    document.add_paragraph("Appendix content belongs here.")
    buffer = io.BytesIO()
    document.save(buffer)

    extracted = extract(buffer.getvalue(), "SOP-QA-001 v11 SOP Writing Numberering Archiving.docx")
    chunks = chunk_document(extracted)
    semantic = [chunk for chunk in chunks if chunk.is_semantic]
    approval_chunk = next(chunk for chunk in chunks if not chunk.is_semantic)

    assert approval_chunk.blocks[0]["type"] == "approval_signature"
    assert approval_chunk.blocks[0]["structured"]["rows"][0] == ["Name", "Signature", "Date"]
    paths = [tuple(chunk.heading_path) for chunk in semantic]
    assert all(path for path in paths)
    assert len({chunk.section_id for chunk in semantic if chunk.heading_path[-1] == "1. References"}) == 1
    assert len({chunk.section_id for chunk in semantic if chunk.heading_path[-1] == "5.1 General"}) == 1
    assert len({chunk.section_id for chunk in semantic if chunk.heading_path[-1] == "5.2 Numbering of SOPs and WIs"}) == 1
    assert len({chunk.section_id for chunk in semantic if chunk.heading_path[-1] == "6. SOP Changes"}) == 1
    appendix = next(chunk for chunk in semantic if "Appendix content" in chunk.text)
    assert appendix.heading_path == ["Appendix I", "SOP and WI Content"]
    assert "6. SOP Changes" not in appendix.heading_path
    assert appendix.blocks[0]["type"] == "appendix_revision"
    table_chunk = next(chunk for chunk in semantic if "SOP-XX-000" in chunk.text)
    assert table_chunk.blocks[0]["type"] == "table"
