"""Resume embedding for documents with existing chunks but NULL vectors."""

from __future__ import annotations

import io

import docx
from sqlalchemy import select

from app.core.config import get_settings
from app.integrations.llm.factory import reset_llm_provider
from app.modules.documents.models import DocumentChunk
from tests.test_ai_routing import FakeGemini, FakeLocal, FakeRemote, _router


def _make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("1. Purpose", level=1)
    document.add_paragraph("Qualification and training requirements.")
    document.add_heading("2. Scope", level=2)
    document.add_paragraph("Applies to all manufacturing staff.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_fallback_mode_persists_local_768_on_ingest(
    client, db_session, auth_headers, company, monkeypatch
):
    monkeypatch.setenv("AI_FEATURES_ENABLED", "true")
    monkeypatch.setenv("AI_ROUTING_MODE", "fallback")
    get_settings.cache_clear()
    reset_llm_provider()

    remote = FakeRemote(llm_up=False, embed_up=False)
    local = FakeLocal()
    provider = _router(mode="fallback", remote=remote, gemini=FakeGemini(), local=local)
    provider._compat.compatible = False
    monkeypatch.setattr(
        "app.integrations.llm.factory.get_llm_provider",
        lambda: provider,
    )

    try:
        response = client.post(
            f"/api/v1/companies/{company['id']}/documents",
            headers=auth_headers,
            files={"file": ("fallback-embed.docx", _make_docx())},
        )
        assert response.status_code == 201, response.text
        detail = response.json()
        assert detail["embeddedChunkCount"] == detail["semanticChunkCount"]
        assert detail["embeddedChunkCount"] > 0

        chunk = db_session.scalar(
            select(DocumentChunk).where(
                DocumentChunk.document_id == detail["id"],
                DocumentChunk.embedding.is_not(None),
            )
        )
        assert chunk is not None
        assert len(list(chunk.embedding)) == 768
        assert chunk.embedding_model == local.model_id
        assert local.embed_calls >= 1
    finally:
        get_settings.cache_clear()
        monkeypatch.undo()
        reset_llm_provider()


def test_retry_embeddings_fills_null_vectors(
    client, db_session, auth_headers, company, monkeypatch
):
    monkeypatch.setenv("AI_FEATURES_ENABLED", "true")
    monkeypatch.setenv("AI_ROUTING_MODE", "fallback")
    get_settings.cache_clear()
    reset_llm_provider()

    remote = FakeRemote(llm_up=False, embed_up=False)
    local = FakeLocal()
    provider = _router(mode="fallback", remote=remote, gemini=FakeGemini(), local=local)
    provider._compat.compatible = False
    monkeypatch.setattr(
        "app.integrations.llm.factory.get_llm_provider",
        lambda: provider,
    )

    try:
        upload = client.post(
            f"/api/v1/companies/{company['id']}/documents",
            headers=auth_headers,
            files={"file": ("retry-me.docx", _make_docx())},
        )
        assert upload.status_code == 201, upload.text
        document_id = upload.json()["id"]

        chunks = list(
            db_session.scalars(
                select(DocumentChunk).where(DocumentChunk.document_id == document_id)
            ).all()
        )
        assert chunks
        for chunk in chunks:
            chunk.embedding = None
            chunk.embedding_model = None
        db_session.flush()

        semantic_count = sum(1 for c in chunks if c.extra.get("isSemantic", True))
        assert semantic_count >= 1

        response = client.post(
            f"/api/v1/companies/{company['id']}/documents/{document_id}/embeddings",
            headers=auth_headers,
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["embeddedChunkCount"] == body["semanticChunkCount"]
        assert body["embeddedChunkCount"] == semantic_count

        refreshed = list(
            db_session.scalars(
                select(DocumentChunk).where(DocumentChunk.document_id == document_id)
            ).all()
        )
        embedded = [
            c for c in refreshed if c.extra.get("isSemantic", True) and c.embedding is not None
        ]
        assert len(embedded) == semantic_count
        assert all(len(list(c.embedding)) == 768 for c in embedded)
        assert local.embed_calls >= 1
    finally:
        get_settings.cache_clear()
        monkeypatch.undo()
        reset_llm_provider()
