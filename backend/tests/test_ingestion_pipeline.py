"""
Regression fixtures for layout-aware SOP ingestion.

Covers intentionally different structures that must normalize into usable
semantic blocks with provenance.
"""

from __future__ import annotations

import io

import docx
import pytest
from docx.shared import Pt
from pypdf import PdfWriter

from app.processing.chunker import chunk_document
from app.processing.extractor import extract
from app.processing.ocr import OcrPageResult, reset_ocr_provider, set_ocr_provider
from app.processing.pipeline import run_pipeline
from app.processing.schema import NormalizedBlock


def _save_docx(document: docx.Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _styled_docx() -> bytes:
    """A. DOCX with Heading styles."""
    document = docx.Document()
    document.add_heading("SOP-014 Lieferantenqualifizierung", level=1)
    document.add_heading("1. Zweck", level=2)
    document.add_paragraph("Diese SOP beschreibt die Qualifizierung externer Lieferanten.")
    document.add_heading("2. Verantwortlichkeiten", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Rolle"
    table.cell(0, 1).text = "Verantwortung"
    table.cell(1, 0).text = "QA Manager"
    table.cell(1, 1).text = "Freigabe der Qualifizierung"
    document.add_heading("3. Ablauf", level=2)
    for step in ["Lieferantenauswahl", "Auditplanung", "Bewertung"]:
        document.add_paragraph(step, style="List Number")
    return _save_docx(document)


def _visual_heading_docx() -> bytes:
    """B. Visually formatted headings, no Heading styles."""
    document = docx.Document()
    title = document.add_paragraph()
    run = title.add_run("SOP-VIS-001 Visual Heading Document")
    run.bold = True
    run.font.size = Pt(18)
    h1 = document.add_paragraph()
    run = h1.add_run("1. Objective")
    run.bold = True
    run.font.size = Pt(14)
    document.add_paragraph("Describe the controlled process without style-based headings.")
    h2 = document.add_paragraph()
    run = h2.add_run("2. Method")
    run.bold = True
    run.font.size = Pt(14)
    document.add_paragraph("Execute verification steps in sequence.")
    return _save_docx(document)


def _numbered_sop_docx() -> bytes:
    """C. Numbered SOP without relying on Heading 1/2 styles alone."""
    document = docx.Document()
    document.add_paragraph("SOP-NUM-003 Numbered Control Procedure")
    document.add_paragraph("1. Purpose")
    document.add_paragraph("Define numbering behaviour for controlled documents.")
    document.add_paragraph("1.1 Applicability")
    document.add_paragraph("Applies to manufacturing and QC areas.")
    document.add_paragraph("2. Responsibilities")
    document.add_paragraph("QA owns the controlled copy.")
    return _save_docx(document)


def _table_dominated_docx() -> bytes:
    """D. SOP dominated by tables."""
    document = docx.Document()
    document.add_heading("SOP-TBL-010 Tabular Process", level=1)
    document.add_heading("1. Role Matrix", level=2)
    table = document.add_table(rows=4, cols=3)
    headers = ("Role", "Activity", "Record")
    for i, value in enumerate(headers):
        table.cell(0, i).text = value
    rows = [
        ("Operator", "Execute step", "Batch record"),
        ("QA", "Review", "Checklist"),
        ("QP", "Release", "Certificate"),
    ]
    for r, values in enumerate(rows, start=1):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    document.add_heading("2. Notes", level=2)
    document.add_paragraph("Table content is authoritative.")
    return _save_docx(document)


def _approval_revision_docx() -> bytes:
    """H. Approval signature block + revision history."""
    document = docx.Document()
    approval = document.add_table(rows=2, cols=3)
    for index, value in enumerate(("Name", "Signature", "Date")):
        approval.cell(0, index).text = value
        approval.cell(1, index).text = "QA Approver"
    document.add_heading("1. Procedure", level=1)
    document.add_paragraph("Perform the activity as written.")
    document.add_heading("Revision History", level=1)
    history = document.add_table(rows=3, cols=3)
    for index, value in enumerate(("Rev", "Date", "Description")):
        history.cell(0, index).text = value
    history.cell(1, 0).text = "01"
    history.cell(1, 1).text = "2024-01-01"
    history.cell(1, 2).text = "Initial release"
    history.cell(2, 0).text = "02"
    history.cell(2, 1).text = "2025-06-01"
    history.cell(2, 2).text = "Clarified responsibilities"
    return _save_docx(document)


def _unusual_sections_docx() -> bytes:
    """I. Unusual section names/order — no PURPOSE/SCOPE."""
    document = docx.Document()
    document.add_heading("SOP-ZZ-900 Odd Structure", level=1)
    document.add_heading("Alpha Gate", level=2)
    document.add_paragraph("Entry criteria for the unusual flow.")
    document.add_heading("Zebra Checks", level=2)
    document.add_paragraph("Verification happens before documentation.")
    document.add_heading("Moonlight Records", level=2)
    document.add_paragraph("Store evidence in the controlled vault.")
    return _save_docx(document)


def _digital_pdf() -> bytes:
    """E. Digital structured PDF with text layer."""
    pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    _, height = A4
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 72, "SOP-PDF-001 Digital Structured Document")
    c.setFont("Helvetica-Bold", 12)
    c.drawString(72, height - 110, "1. Purpose")
    c.setFont("Helvetica", 11)
    c.drawString(72, height - 130, "This PDF has a real text layer for extraction.")
    c.setFont("Helvetica-Bold", 12)
    c.drawString(72, height - 160, "2. Procedure")
    c.setFont("Helvetica", 11)
    c.drawString(72, height - 180, "Follow the numbered steps below.")
    c.drawString(72, height - 200, "1) Prepare materials")
    c.drawString(72, height - 220, "2) Execute controls")
    c.showPage()
    c.save()
    return buffer.getvalue()


def _repeated_header_pdf() -> bytes:
    """G. Repeated header/footer furniture across pages."""
    pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    _, height = A4
    for page in range(1, 4):
        c.setFont("Helvetica", 9)
        c.drawString(72, height - 36, "CONFIDENTIAL — SOP-FURN-001 — Controlled Copy")
        c.drawString(72, 36, f"Page {page} of 3 — Doc Ctrl")
        c.setFont("Helvetica-Bold", 14)
        c.drawString(72, height - 90, f"{page}. Section Title {page}")
        c.setFont("Helvetica", 11)
        c.drawString(72, height - 120, f"Semantic body content for page {page}.")
        c.showPage()
    c.save()
    return buffer.getvalue()


def _scanned_pdf() -> bytes:
    """F. Image-only / blank pages → scanned classification."""
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@pytest.fixture(autouse=True)
def _reset_ocr():
    reset_ocr_provider()
    yield
    reset_ocr_provider()


class TestStyledDocx:
    def test_a_heading_styles_normalize(self):
        result = extract(_styled_docx(), "a-styled.docx")
        kinds = [b.kind for b in result.blocks]
        assert kinds.count("heading") >= 4
        assert kinds.count("table") == 1
        assert kinds.count("list") >= 3
        assert result.qa_status in ("PASS", "WARNING")
        chunks = chunk_document(result)
        assert chunks and all(c.heading_path for c in chunks if c.is_semantic)


class TestVisualHeadings:
    def test_b_visual_headings_without_styles(self):
        result = extract(_visual_heading_docx(), "b-visual.docx")
        headings = [b for b in result.blocks if b.kind == "heading"]
        assert len(headings) >= 2
        assert any("Objective" in h.text for h in headings)
        assert any("Method" in h.text for h in headings)
        assert any(
            "font_size" in (b.structured.get("signals") or [])
            or "bold" in (b.structured.get("signals") or [])
            for b in headings
        )


class TestNumberedSop:
    def test_c_numbered_hierarchy(self):
        result = extract(_numbered_sop_docx(), "c-numbered.docx")
        headings = [b for b in result.blocks if b.kind == "heading"]
        assert any(b.text.startswith("1.") for b in headings)
        assert any(b.text.startswith("1.1") for b in headings)
        chunks = chunk_document(result)
        semantic = [c for c in chunks if c.is_semantic]
        assert any(len(c.heading_path) >= 2 for c in semantic)


class TestTableDominated:
    def test_d_tables_stay_atomic(self):
        result = extract(_table_dominated_docx(), "d-tables.docx")
        table = next(b for b in result.blocks if b.kind == "table")
        assert table.structured.get("rows")
        assert "QA" in table.text
        chunks = chunk_document(result)
        table_chunks = [c for c in chunks if any(b["type"] == "table" for b in c.blocks)]
        assert len(table_chunks) == 1
        assert "Operator" in table_chunks[0].text


class TestDigitalPdf:
    def test_e_digital_pdf_structure(self):
        result = extract(_digital_pdf(), "e-digital.pdf")
        assert result.source_format == "pdf"
        assert result.pdf_kind == "digital"
        assert not result.is_probably_scanned
        assert any(b.kind == "heading" for b in result.blocks)
        assert any("Purpose" in b.text for b in result.blocks)
        assert result.qa_status in ("PASS", "WARNING")


class TestScannedMixedPdf:
    def test_f_scanned_flagged_and_ocr_provider_invoked(self):
        class RecordingOcr:
            name = "recording"

            def __init__(self):
                self.calls = 0

            def ocr_pages(self, pages):
                self.calls += 1
                return [
                    OcrPageResult(
                        page_number=page.page_number,
                        blocks=[],
                        engine=self.name,
                        warnings=[f"Page {page.page_number} needs production OCR."],
                    )
                    for page in pages
                ]

        provider = RecordingOcr()
        set_ocr_provider(provider)
        result = extract(_scanned_pdf(), "f-scanned.pdf")
        assert result.is_probably_scanned
        assert result.pdf_kind == "scanned"
        assert result.qa_status == "FAILED"
        assert provider.calls == 1
        assert any("OCR" in w for w in result.warnings)
        assert result.blocks == []


class TestRepeatedFurniture:
    def test_g_repeated_headers_marked_non_semantic(self):
        result = extract(_repeated_header_pdf(), "g-furniture.pdf")
        furniture = [b for b in result.blocks if b.kind == "header_footer"]
        assert furniture
        assert all(not b.semantic for b in furniture)
        chunks = chunk_document(result)
        semantic = [c for c in chunks if c.is_semantic]
        assert semantic
        assert not any("CONFIDENTIAL" in c.text for c in semantic)


class TestApprovalRevision:
    def test_h_approval_and_revision_history(self):
        result = extract(_approval_revision_docx(), "h-approval.docx")
        assert any(b.kind == "approval_signature" and not b.semantic for b in result.blocks)
        assert any(b.kind == "revision_history" for b in result.blocks)
        chunks = chunk_document(result)
        assert any(not c.is_semantic for c in chunks)
        assert any(
            any(b["type"] == "revision_history" for b in c.blocks) for c in chunks
        )


class TestUnusualSections:
    def test_i_unusual_names_still_structure(self):
        result = extract(_unusual_sections_docx(), "i-unusual.docx")
        headings = [b.text for b in result.blocks if b.kind == "heading"]
        assert "Alpha Gate" in headings
        assert "Zebra Checks" in headings
        chunks = chunk_document(result)
        paths = [" > ".join(c.heading_path) for c in chunks if c.is_semantic]
        assert any("Moonlight Records" in path for path in paths)


def test_pipeline_exposes_normalized_schema():
    normalized = run_pipeline(_styled_docx(), "pipe.docx")
    assert normalized.blocks
    assert normalized.extraction_quality.status in ("PASS", "WARNING", "FAILED")
    assert all(isinstance(b, NormalizedBlock) for b in normalized.blocks)
    assert all(b.id and b.order >= 0 for b in normalized.blocks)
