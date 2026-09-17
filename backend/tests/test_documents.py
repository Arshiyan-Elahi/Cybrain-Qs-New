import io

import docx
import pytest
from datetime import UTC, datetime
from sqlalchemy import select

from app.core.config import Settings
from app.modules.documents.dependencies import get_ingestion_service
from app.modules.documents.models import Document, DocumentChunk
from app.modules.knowledge.models import KnowledgeObject
from app.modules.auth.models import UserCompanyAccess
from app.shared.enums import KnowledgeStatus, KnowledgeTier


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("SOP-021 Reinigungsvalidierung", level=1)
    document.add_heading("1. Zweck", level=2)
    document.add_paragraph("Diese SOP beschreibt die Validierung von Reinigungsverfahren.")
    document.add_heading("2. Verantwortlichkeiten", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Rolle"
    table.cell(0, 1).text = "Verantwortung"
    table.cell(1, 0).text = "QA Manager"
    table.cell(1, 1).text = "Freigabe"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(client, company_id, headers, filename="SOP-021.docx", content=None):
    return client.post(
        f"/api/v1/companies/{company_id}/documents",
        headers=headers,
        files={"file": (filename, content if content is not None else make_docx())},
    )


@pytest.fixture
def uploaded(client, auth_headers, company):
    response = upload(client, company["id"], auth_headers)
    assert response.status_code == 201, response.text
    return response.json()


class TestUploadWithoutAI:
    """Ingestion must work with no inference server running."""

    def test_document_is_processed(self, uploaded):
        assert uploaded["status"] == "processed"
        assert uploaded["chunkCount"] > 0

    def test_no_embeddings_while_ai_disabled(self, uploaded):
        assert uploaded["embeddedChunkCount"] == 0
        assert not any("Embedding" in w for w in uploaded["warnings"])

    def test_structure_is_preserved(self, client, auth_headers, company, uploaded):
        chunks = client.get(
            f"/api/v1/companies/{company['id']}/documents/{uploaded['id']}/chunks",
            headers=auth_headers,
        ).json()
        assert all(c["headingPath"] for c in chunks)
        # A table must survive as one chunk, never split mid-row.
        assert any("QA Manager" in c["text"] for c in chunks)


def test_ingestion_provider_is_injected_only_when_ai_enabled(monkeypatch, db_session):
    calls = 0
    provider = object()

    def fake_provider():
        nonlocal calls
        calls += 1
        return provider

    monkeypatch.setattr("app.integrations.llm.factory.get_llm_provider", fake_provider)

    without_ai = get_ingestion_service(db_session, Settings(ai_features_enabled=False))
    assert without_ai.llm is None
    assert calls == 0

    with_ai = get_ingestion_service(db_session, Settings(ai_features_enabled=True))
    assert with_ai.llm is provider
    assert calls == 1


class TestUploadValidation:
    def test_unsupported_type_rejected(self, client, auth_headers, company):
        response = upload(client, company["id"], auth_headers, "notes.txt", b"just text")
        assert response.status_code == 422
        assert "Supported" in response.json()["error"]["message"]

    def test_empty_file_rejected(self, client, auth_headers, company):
        response = upload(client, company["id"], auth_headers, "empty.pdf", b"")
        assert response.status_code == 422

    def test_corrupt_pdf_recorded_with_reason(self, client, auth_headers, company):
        response = upload(client, company["id"], auth_headers, "broken.pdf", b"%PDF-1.4 garbage")
        assert response.status_code == 201
        body = response.json()
        # Stored as failed WITH an explanation, never silently dropped.
        assert body["status"] == "failed"
        assert body["warnings"]

    def test_multiple_uploads_are_independent_and_listed(
        self, client, auth_headers, company
    ):
        first = upload(client, company["id"], auth_headers, "first.docx")
        broken = upload(
            client,
            company["id"],
            auth_headers,
            "broken.pdf",
            b"%PDF-1.4 garbage",
        )
        second = upload(client, company["id"], auth_headers, "second.docx")

        assert first.status_code == 201
        assert first.json()["status"] == "processed"
        assert broken.status_code == 201
        assert broken.json()["status"] == "failed"
        assert second.status_code == 201
        assert second.json()["status"] == "processed"

        listed = client.get(
            f"/api/v1/companies/{company['id']}/documents", headers=auth_headers
        )
        assert listed.status_code == 200
        by_name = {item["filename"]: item["status"] for item in listed.json()["items"]}
        assert by_name == {
            "first.docx": "processed",
            "broken.pdf": "failed",
            "second.docx": "processed",
        }


class TestStats:
    def test_counters(self, client, auth_headers, company, uploaded):
        stats = client.get(f"/api/v1/companies/{company['id']}/stats", headers=auth_headers).json()
        assert stats["documentCount"] == 1
        assert stats["processedCount"] == 1
        assert stats["chunkCount"] > 0
        assert stats["totalBytes"] > 0

    def test_reports_ai_disabled(self, client, auth_headers, company):
        stats = client.get(f"/api/v1/companies/{company['id']}/stats", headers=auth_headers).json()
        assert stats["aiFeaturesEnabled"] is False
        assert stats["embeddedChunkCount"] == 0


class TestDocumentIsolation:
    def test_other_user_cannot_list(self, client, company, uploaded, user_factory):
        other = user_factory()
        assert (
            client.get(f"/api/v1/companies/{company['id']}/documents", headers=other).status_code
            == 404
        )


class TestDocumentDeletion:
    def test_deletes_chunks_embeddings_and_proposed_knowledge_only_for_target(
        self, client, db_session, auth_headers, company
    ):
        target = upload(client, company["id"], auth_headers, "delete-me.docx").json()
        keep = upload(client, company["id"], auth_headers, "keep-me.docx").json()
        chunk = db_session.scalar(select(DocumentChunk).where(DocumentChunk.document_id == target["id"]))
        chunk.embedding = [0.0] * 768
        access = db_session.scalar(select(UserCompanyAccess).where(UserCompanyAccess.company_id == company["id"]))
        db_session.add(KnowledgeObject(
            company_id=company["id"], type="terminology", tier=KnowledgeTier.COMPANY,
            status=KnowledgeStatus.PROPOSED, label="Delete with source", payload={},
            source_document_id=target["id"], source_chunk_id=chunk.id,
            source_location=chunk.location, extraction_method="test",
            extracted_at=datetime.now(UTC),
        ))
        db_session.flush()

        response = client.delete(f"/api/v1/companies/{company['id']}/documents/{target['id']}", headers=auth_headers)
        assert response.status_code == 204
        assert db_session.get(Document, target["id"]) is None
        assert db_session.scalar(select(DocumentChunk.id).where(DocumentChunk.document_id == target["id"])) is None
        assert db_session.scalar(select(KnowledgeObject.id).where(KnowledgeObject.source_document_id == target["id"])) is None
        assert db_session.get(Document, keep["id"]) is not None
        assert client.delete(f"/api/v1/companies/{company['id']}/documents/{target['id']}", headers=auth_headers).status_code == 204

    def test_verified_provenance_blocks_delete(self, client, db_session, auth_headers, company):
        target = upload(client, company["id"], auth_headers, "verified-source.docx").json()
        chunk = db_session.scalar(select(DocumentChunk).where(DocumentChunk.document_id == target["id"]))
        access = db_session.scalar(select(UserCompanyAccess).where(UserCompanyAccess.company_id == company["id"]))
        db_session.add(KnowledgeObject(
            company_id=company["id"], type="business_rule", tier=KnowledgeTier.COMPANY,
            status=KnowledgeStatus.VERIFIED, label="Verified", payload={},
            source_document_id=target["id"], source_chunk_id=chunk.id,
            source_location=chunk.location, extraction_method="test",
            extracted_at=datetime.now(UTC), verified_by=access.user_id, verified_at=datetime.now(UTC),
        ))
        db_session.flush()
        response = client.delete(f"/api/v1/companies/{company['id']}/documents/{target['id']}", headers=auth_headers)
        assert response.status_code == 409
        body = response.json()["error"]
        assert "verified company knowledge" in body["message"]
        assert body["details"]["verifiedDependencyCount"] == 1
        assert body["details"]["reason"] == "verified_knowledge_dependency"

    def test_other_user_cannot_read(self, client, company, uploaded, user_factory):
        other = user_factory()
        assert (
            client.get(
                f"/api/v1/companies/{company['id']}/documents/{uploaded['id']}", headers=other
            ).status_code
            == 404
        )

    def test_other_user_cannot_upload(self, client, company, user_factory):
        other = user_factory()
        assert upload(client, company["id"], other).status_code == 404

    def test_other_user_cannot_read_stats(self, client, company, user_factory):
        other = user_factory()
        assert (
            client.get(f"/api/v1/companies/{company['id']}/stats", headers=other).status_code == 404
        )
