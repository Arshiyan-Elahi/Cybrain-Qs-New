"""Company create with onboarding profile persistence and idempotency."""

import uuid


def _onboarding_payload(**overrides):
    base = {
        "startOptionId": "nothing",
        "documentPathId": "continue-without",
        "structurePreferenceId": "standard",
        "documentStructureNotes": "Keep sections short",
        "toneId": "neutral",
        "formalityId": "formal",
        "personId": "third",
        "writingNotes": "Prefer clear verbs",
        "terminologyText": "batch record",
        "preferExistingTerms": True,
        "departmentsText": "QA, Manufacturing",
        "rolesText": "QA Manager",
        "processesText": "Deviation handling",
        "workflowNotes": "Two-step review",
        "templatePreferenceId": "gmp-default",
        "layoutNotes": "Numbered headings",
        "formsText": "Change control form",
        "businessRulesText": "No silent overwrites",
        "relationshipsText": "SOP links to forms",
        "bestPracticesText": "Cite sources",
        "aiAssistLevelId": "balanced",
        "requireHumanVerification": True,
        "qualityNotes": "Human gate required",
        "intendedDocumentNames": [],
        "intendedTemplateNames": [],
    }
    base.update(overrides)
    return base


def test_create_persists_onboarding_profile(client, auth_headers):
    request_id = str(uuid.uuid4())
    response = client.post(
        "/api/v1/companies",
        headers=auth_headers,
        json={
            "name": "Onboarded GmbH",
            "industryKey": "pharma",
            "locationKey": "vienna-at",
            "primaryLanguageKey": "en",
            "regulationIds": ["eu-gmp"],
            "creationRequestId": request_id,
            "onboarding": _onboarding_payload(
                startOptionId="existing-sops",
                documentPathId="upload-sops",
                intendedDocumentNames=["sop-a.pdf"],
            ),
        },
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["name"] == "Onboarded GmbH"
    assert body["creationRequestId"] == request_id
    assert body["onboarding"]["startOptionId"] == "existing-sops"
    assert body["onboarding"]["documentPathId"] == "upload-sops"
    assert body["onboarding"]["terminologyText"] == "batch record"
    assert body["onboarding"]["intendedDocumentNames"] == ["sop-a.pdf"]
    assert body["onboarding"]["requireHumanVerification"] is True

    fetched = client.get(f"/api/v1/companies/{body['id']}", headers=auth_headers)
    assert fetched.status_code == 200
    assert fetched.json()["onboarding"]["writingNotes"] == "Prefer clear verbs"


def test_create_idempotent_by_creation_request_id(client, auth_headers):
    request_id = str(uuid.uuid4())
    payload = {
        "name": "Once Only AG",
        "industryKey": "biotech",
        "locationKey": "berlin-de",
        "creationRequestId": request_id,
        "onboarding": _onboarding_payload(),
    }
    first = client.post("/api/v1/companies", headers=auth_headers, json=payload)
    second = client.post("/api/v1/companies", headers=auth_headers, json=payload)
    assert first.status_code == 201
    assert second.status_code == 201
    assert first.json()["id"] == second.json()["id"]
    assert client.get("/api/v1/companies", headers=auth_headers).json()["total"] == 1


def test_nothing_path_creates_without_documents(client, auth_headers):
    response = client.post(
        "/api/v1/companies",
        headers=auth_headers,
        json={
            "name": "Blank Start Ltd",
            "industryKey": "medtech",
            "locationKey": "zurich-ch",
            "creationRequestId": str(uuid.uuid4()),
            "onboarding": _onboarding_payload(
                startOptionId="nothing",
                documentPathId="continue-without",
            ),
        },
    )
    assert response.status_code == 201
    company_id = response.json()["id"]
    docs = client.get(f"/api/v1/companies/{company_id}/documents", headers=auth_headers)
    assert docs.status_code == 200
    assert docs.json()["total"] == 0


def test_template_upload_persists_kind(client, auth_headers):
    import io

    import docx

    document = docx.Document()
    document.add_heading("SOP Template", level=1)
    document.add_paragraph("Header block")
    buffer = io.BytesIO()
    document.save(buffer)

    created = client.post(
        "/api/v1/companies",
        headers=auth_headers,
        json={
            "name": "Template Path Co",
            "industryKey": "pharma",
            "locationKey": "vienna-at",
            "creationRequestId": str(uuid.uuid4()),
            "onboarding": _onboarding_payload(
                startOptionId="template",
                documentPathId="upload-template",
                intendedTemplateNames=["template.docx"],
            ),
        },
    )
    assert created.status_code == 201
    company_id = created.json()["id"]
    uploaded = client.post(
        f"/api/v1/companies/{company_id}/documents",
        headers=auth_headers,
        params={"kind": "template"},
        files={"file": ("template.docx", buffer.getvalue())},
    )
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["kind"] == "template"
    assert uploaded.json()["filename"] == "template.docx"


class TestOnboardingTenantIsolation:
    def test_other_user_cannot_read_onboarding(self, client, auth_headers, user_factory):
        created = client.post(
            "/api/v1/companies",
            headers=auth_headers,
            json={
                "name": "Private Profile Co",
                "industryKey": "pharma",
                "locationKey": "vienna-at",
                "creationRequestId": str(uuid.uuid4()),
                "onboarding": _onboarding_payload(terminologyText="secret term"),
            },
        )
        assert created.status_code == 201
        company_id = created.json()["id"]
        other = user_factory()
        assert client.get(f"/api/v1/companies/{company_id}", headers=other).status_code == 404
