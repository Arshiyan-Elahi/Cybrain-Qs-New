"""Company permanent deletion — tenant-safe cascade of company-owned data."""

from datetime import UTC, datetime
import uuid

from sqlalchemy import select

from app.modules.auth.models import User, UserCompanyAccess
from app.modules.companies.models import Company, CompanyOnboardingProfile
from app.modules.documents.models import Document, DocumentChunk
from app.modules.knowledge.models import KnowledgeObject, KnowledgeObjectHistory
from app.shared.enums import KnowledgeStatus, KnowledgeTier


def _upload(client, company_id, headers, filename="sop.docx"):
    import io
    import docx

    document = docx.Document()
    document.add_heading("Purpose", level=1)
    document.add_paragraph("Delete coverage text.")
    buffer = io.BytesIO()
    document.save(buffer)
    return client.post(
        f"/api/v1/companies/{company_id}/documents",
        headers=headers,
        files={"file": (filename, buffer.getvalue())},
    )


def test_delete_removes_company_owned_data_and_spares_other_tenant(
    client, db_session, auth_headers, company, user_factory
):
    other_headers = user_factory()
    other = client.post(
        "/api/v1/companies",
        headers=other_headers,
        json={
            "name": "Other Tenant",
            "industryKey": "pharma",
            "locationKey": "vienna-at",
            "regulationIds": ["eu-gmp"],
        },
    ).json()

    target = company
    doc = _upload(client, target["id"], auth_headers, "delete-co.docx").json()
    other_doc = _upload(client, other["id"], other_headers, "keep-co.docx").json()

    chunk = db_session.scalar(select(DocumentChunk).where(DocumentChunk.document_id == doc["id"]))
    chunk.embedding = [0.05] * 768
    other_chunk = db_session.scalar(
        select(DocumentChunk).where(DocumentChunk.document_id == other_doc["id"])
    )
    other_chunk.embedding = [0.09] * 768

    access = db_session.scalar(
        select(UserCompanyAccess).where(UserCompanyAccess.company_id == target["id"])
    )
    assert access is not None and access.role == "owner"
    owner_user_id = access.user_id

    db_session.add(
        KnowledgeObject(
            company_id=target["id"],
            type="terminology",
            tier=KnowledgeTier.COMPANY,
            status=KnowledgeStatus.VERIFIED,
            label="Verified KO",
            payload={"evidence": [{"documentId": doc["id"], "snippet": "x"}]},
            source_kind="uploaded_document",
            source_document_id=doc["id"],
            source_chunk_id=chunk.id,
            source_location=chunk.location,
            extraction_method="test",
            extracted_at=datetime.now(UTC),
            verified_by=owner_user_id,
            verified_at=datetime.now(UTC),
        )
    )
    db_session.add(
        KnowledgeObject(
            company_id=other["id"],
            type="terminology",
            tier=KnowledgeTier.INDUSTRY,
            status=KnowledgeStatus.PROPOSED,
            label="Other industry-tier KO",
            payload={},
            source_kind="ai_extracted",
            source_document_id=other_doc["id"],
            source_chunk_id=other_chunk.id,
            source_location=other_chunk.location,
            extraction_method="test",
            extracted_at=datetime.now(UTC),
        )
    )
    db_session.flush()
    ko = db_session.scalar(
        select(KnowledgeObject).where(KnowledgeObject.company_id == target["id"])
    )
    db_session.add(
        KnowledgeObjectHistory(
            company_id=target["id"],
            knowledge_object_id=ko.id,
            action="confirmed",
            actor_id=owner_user_id,
            label=ko.label,
            status=ko.status,
            source_kind=ko.source_kind,
            version=ko.version,
            evidence_snapshot=[],
            payload_snapshot={},
            detail="fixture",
        )
    )
    db_session.flush()

    if db_session.get(CompanyOnboardingProfile, uuid.UUID(target["id"])) is None:
        db_session.add(
            CompanyOnboardingProfile(
                company_id=target["id"],
                start_option_id="docs",
                document_path_id="upload",
            )
        )
        db_session.flush()

    response = client.delete(f"/api/v1/companies/{target['id']}", headers=auth_headers)
    assert response.status_code == 204

    assert db_session.get(Company, target["id"]) is None
    assert db_session.get(CompanyOnboardingProfile, uuid.UUID(target["id"])) is None
    assert db_session.get(Document, doc["id"]) is None
    assert (
        db_session.scalar(select(DocumentChunk.id).where(DocumentChunk.company_id == target["id"]))
        is None
    )
    assert (
        db_session.scalar(
            select(DocumentChunk.id).where(
                DocumentChunk.company_id == target["id"],
                DocumentChunk.embedding.is_not(None),
            )
        )
        is None
    )
    assert (
        db_session.scalar(
            select(KnowledgeObject.id).where(KnowledgeObject.company_id == target["id"])
        )
        is None
    )
    assert (
        db_session.scalar(
            select(KnowledgeObjectHistory.id).where(
                KnowledgeObjectHistory.company_id == target["id"]
            )
        )
        is None
    )
    assert (
        db_session.scalar(
            select(UserCompanyAccess.id).where(UserCompanyAccess.company_id == target["id"])
        )
        is None
    )

    # Other tenant untouched.
    assert db_session.get(Company, other["id"]) is not None
    assert db_session.get(Document, other_doc["id"]) is not None
    assert (
        db_session.scalar(
            select(DocumentChunk.id).where(DocumentChunk.document_id == other_doc["id"])
        )
        is not None
    )
    assert (
        db_session.scalar(
            select(KnowledgeObject.id).where(KnowledgeObject.company_id == other["id"])
        )
        is not None
    )

    # User account remains.
    assert db_session.get(User, owner_user_id) is not None
    assert client.get("/api/v1/auth/me", headers=auth_headers).status_code == 200


def test_member_cannot_delete_company(client, db_session, auth_headers, company, user_factory):
    member_headers = user_factory()
    member_me = client.get("/api/v1/auth/me", headers=member_headers).json()
    db_session.add(
        UserCompanyAccess(
            user_id=member_me["id"],
            company_id=company["id"],
            role="member",
        )
    )
    db_session.flush()

    response = client.delete(f"/api/v1/companies/{company['id']}", headers=member_headers)
    assert response.status_code == 403
    assert response.json()["error"]["details"]["reason"] == "company_owner_required"
    assert db_session.get(Company, company["id"]) is not None


def test_other_user_without_access_cannot_delete(client, company, user_factory):
    other = user_factory()
    assert client.delete(f"/api/v1/companies/{company['id']}", headers=other).status_code == 404


def test_delete_cancels_in_flight_document_operations(client, auth_headers, company):
    from app.modules.documents.cancellation import operations

    doc = _upload(client, company["id"], auth_headers, "cancel-co.docx").json()
    op_id = uuid.uuid4()
    operations.bind_document(op_id, uuid.UUID(doc["id"]))
    assert client.delete(f"/api/v1/companies/{company['id']}", headers=auth_headers).status_code == 204
    assert op_id not in operations._states
